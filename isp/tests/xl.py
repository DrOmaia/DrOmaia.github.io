import openpyxl, sys
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb.worksheets[0]
print('sheets', wb.sheetnames)
for r in range(4, 9): print(r, [ws.cell(r, c).value for c in range(1, 14)])
# check table rows of interest
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    a = row[0].value
    if a in ('IS450-444', 'IS321-428', 'IS201-498', 'IS492-547'): print('check', [c.value for c in row[:9]])
ir = wb['Is Reg']
hdr = [c.value for c in ir[2]]; print('isreg hdr', hdr)
for row in ir.iter_rows(min_row=3):
    v = [c.value for c in row]
    if v[1] in ('IS321', 'IS450') or v[0] in (428, 444): print('isreg', v[:12])
